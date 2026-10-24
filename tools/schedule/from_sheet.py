"""'2026 직원 스케쥴' 구글 시트(xlsx로 받은 파일)의 월 탭 → tools/schedule/YYYY-MM.csv

  python3 tools/schedule/from_sheet.py <받은파일.xlsx> 2026 11

병원이 매달 카카오·인스타용 일정 이미지에 쓰는 규칙을 그대로 따른다.
  09:00 · 10:00 출근            → 오전
  11:00 · 12:00 출근, 당직 출근  → 오후
  시간 없이 색만 칠한 칸(휴무·연차·당직 다음날·예비군 등) → 휴무
수의사 칸(김율현·김평·김세림·김종완)만 읽는다. 테크니션·야간당직 칸은 읽지 않는다.
결과 파일은 원장님 확인(승인) 뒤에만 커밋한다.
"""
import calendar
import datetime
import pathlib
import re
import sys

import openpyxl

OUT = pathlib.Path(__file__).resolve().parent
ORDER = ['김평', '김율현', '김세림', '김종완']   # 페이지·이미지 표시 순서
WARN = []


def label(v):
    if v is None:
        return '휴무'
    if isinstance(v, datetime.time):
        h = v.hour
    else:
        s = str(v)
        m = re.search(r'(\d{1,2}):(\d{2})', s)
        if not m:
            return '휴무'          # '예비군' 같은 글자만 있는 칸
        h = int(m.group(1))
        if '당직' in s:
            return '오후'
        if not re.fullmatch(r'\s*\d{1,2}:\d{2}\s*', s):
            WARN.append(s)         # '반차17:30' 같은 칸은 사람이 확인
    return '오전' if h <= 10 else '오후'


def main(xlsx, year, month):
    ws = openpyxl.load_workbook(xlsx)[f'{month}월']
    head_row = None
    for r in range(1, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, 10)]
        if '날짜' in vals and '김평' in vals:
            head_row = r
            cols = {v: i + 1 for i, v in enumerate(vals) if v in ORDER}
            date_col = vals.index('날짜') + 1
            break
    if head_row is None or len(cols) != 4:
        sys.exit('표 머리(날짜·김평·김율현·김세림·김종완)를 찾지 못했습니다.')
    days = calendar.monthrange(year, month)[1]
    rows = {}
    for r in range(head_row + 1, head_row + 40):
        d = ws.cell(r, date_col).value
        m = re.fullmatch(r'(\d{1,2})일', str(d or '').strip())
        if not m:
            continue
        rows[int(m.group(1))] = [label(ws.cell(r, cols[n]).value) for n in ORDER]
    missing = [d for d in range(1, days + 1) if d not in rows]
    if missing:
        sys.exit(f'{month}월 시트에 없는 날짜: {missing}')
    out = OUT / f'{year}-{month:02d}.csv'
    lines = [f'# {year}년 {month}월 의료진 진료일정',
             f'# 자료: 2026 직원 스케쥴 시트 {month}월 탭 (오전=9·10시 출근, 오후=11·12시·당직 출근, 그 외=휴무)',
             '날짜,' + ','.join(ORDER)]
    lines += [f'{d},' + ','.join(rows[d]) for d in range(1, days + 1)]
    out.write_text('\n'.join(lines) + '\n', encoding='utf-8')
    print(out)
    for w in WARN:
        print(f'  확인 필요: "{w}" 칸은 규칙과 다를 수 있습니다. CSV를 직접 확인하세요.')


if __name__ == '__main__':
    main(sys.argv[1], int(sys.argv[2]), int(sys.argv[3]))
